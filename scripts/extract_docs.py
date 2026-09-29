"""Extract text from downloaded PDFs/DOCX into documents/<site>/text/<name>.txt.

Usage: python scripts/extract_docs.py <site> [<site> ...]
Prints character counts; a near-zero count means the file is a scan and needs OCR/manual reading.
"""
import sys
from pathlib import Path

from pypdf import PdfReader

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parent.parent


def pdf_text(p):
    return "\n\n".join((pg.extract_text() or "") for pg in PdfReader(p).pages)


def docx_text(p):
    import docx
    d = docx.Document(p)
    parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


for site in sys.argv[1:]:
    folder = ROOT / "documents" / site
    out = folder / "text"
    out.mkdir(exist_ok=True)
    for f in sorted(folder.iterdir()):
        if f.suffix.lower() not in (".pdf", ".docx"):
            continue
        try:
            txt = pdf_text(f) if f.suffix.lower() == ".pdf" else docx_text(f)
        except Exception as e:
            txt = ""
            print(f"ERR {f.name}: {e}")
        (out / f"{f.stem}.txt").write_text(txt, encoding="utf-8")
        print(f"{len(txt.strip()):>7}  {f.name}")
